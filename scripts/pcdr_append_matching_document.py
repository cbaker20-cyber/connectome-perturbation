"""Append the dated matching audit, preserving the previous complete edition."""
from pathlib import Path
import json
import hashlib
from datetime import datetime, timezone
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
old=root/'exports/pcdr_research_record_20260921_completed/Connectome_Research_Record_2026-09-21_Completed.docx'
out=root/'exports/pcdr_research_record_20260921_matching_audit';out.mkdir(exist_ok=True)
source=root/'docs/pcdr/MATCHING_AUDIT_20260921.md'
doc=Document(old)
doc.paragraphs[4].text=('The 720-trial individual-cell study is complete. One motor effect passed secondary correction. '
    'The original first-20 mode selection failed; an exploratory search selected a recruited support but its sampler found no controls. '
    'The later audit in Chapter 10 found five sets passing unchanged matching thresholds. These optimized examples establish feasibility, '
    'not a random null distribution. No eigen-set lesion comparison has run; the primary hypothesis remains untested. Earlier chapters retain their dated history.')
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('Status snapshot:'):
        paragraph.text='Earlier chapters preserve the 21 September morning snapshot. Chapter 10 adds the matching audit completed later that day; source_manifest.json records this edition.'
    if paragraph.text.startswith('Read Chapters 1'):
        paragraph.text='Read Chapter 10 for the latest matching result and next design. Chapters 1–9 preserve the study, paper notes, eigencircuit guide, code explanations and notebook through the earlier snapshot.'
h=doc.add_heading('10 Matching feasibility audit and next design',level=1);h.paragraph_format.page_break_before=True
lines=source.read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
    line=lines[i]
    if not line or line.startswith('# '):i+=1;continue
    if line.startswith('## '):doc.add_heading(line[3:],level=2)
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            if '---' not in lines[i]:rows.append([c.strip() for c in lines[i].strip('|').split('|')])
            i+=1
        table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Table Grid'
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        table._tbl.tblPr.append(borders)
        for j,value in enumerate(rows[0]):table.rows[0].cells[j].text=value
        for row in rows[1:]:
            for j,value in enumerate(row):table.add_row() if j==0 else None;table.rows[-1].cells[j].text=value
        for cell in table.rows[0].cells:
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E4EDF2');cell._tc.get_or_add_tcPr().append(shd)
            for r in cell.paragraphs[0].runs:r.bold=True
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(4)
                    for run in p.runs:run.font.size=Pt(10)
        continue
    else:doc.add_paragraph(line[2:] if line.startswith('- ') else line,style='List Bullet' if line.startswith('- ') else None)
    i+=1
doc.save(out/'Connectome_Research_Record_2026-09-21_Matching_Audit.docx')
paths=[old,source,Path(__file__),root/'results/pcdr/matching_audit_20260921/summary.json',root/'results/pcdr/matching_audit_20260921/witness/summary.json',root/'results/pcdr/matching_audit_20260921/independent_validation.json']
(out/'source_manifest.json').write_text(json.dumps({'created_utc':datetime.now(timezone.utc).isoformat(),
    'sha256':{str(p.relative_to(root)):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},
    'note':'Cover status updated; original chapters preserved; dated Chapter 10 appended.'},indent=2),encoding='utf-8')
print(out)
