"""Append the dated matching audit, preserving the previous complete edition."""
from pathlib import Path
import json
import hashlib
from datetime import datetime, timezone
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
old=root/'exports/pcdr_research_record_20260922_replication/Connectome_Research_Record_2026-09-22_Replication.docx'
out=root/'exports/pcdr_research_record_20260922_design_audit';out.mkdir(exist_ok=True)
source=root/'docs/pcdr/DESIGN_AUDIT_RECORD_20260922.md'
doc=Document(old)
doc.paragraphs[4].text='The single-cell study and separate fresh-seed replication are complete. Chapter 13 adds the baseline composition audit, restricted matching failure, code explanations and CCR preparation. Comparisons remain conditional on optimized sets; confirmation and actual CCR validation are pending. Earlier chapters preserve the dated research history.'
for p in doc.paragraphs:
    if p.text.startswith('Chapter 12 records'): p.text='Chapter 13 records the latest design audit and CCR preparation on 22 September 2026. Chapter 12 retains every fresh-seed replication result.'
    if p.text.startswith('Read Chapter 12'): p.text='Read Chapter 13 for current status and next steps; Chapter 12 for replication results. Earlier chapters preserve literature, methods, code guides, process notes and prior findings.'
    if '18–21 September 2026' in p.text: p.text=p.text.replace('18–21 September 2026','18–22 September 2026')
h=doc.add_heading('13 Composition audit and CCR preparation',level=1);h.paragraph_format.page_break_before=True
lines=source.read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
    line=lines[i]
    if not line or line.startswith('# '):i+=1;continue
    if line.startswith('## '):doc.add_heading(line[3:],level=2)
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            if '---' not in lines[i]:rows.append([c.strip() for c in lines[i].strip('|').split('|')])
            i+=1
        table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Table Grid'
        header_repeat=OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        table._tbl.tblPr.append(borders)
        for j,value in enumerate(rows[0]):table.rows[0].cells[j].text=value
        for row in rows[1:]:
            for j,value in enumerate(row):table.add_row() if j==0 else None;table.rows[-1].cells[j].text=value
        for cell in table.rows[0].cells:
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E4EDF2');cell._tc.get_or_add_tcPr().append(shd)
            for r in cell.paragraphs[0].runs:r.bold=True
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(4)
                    for run in p.runs:run.font.size=Pt(10)
        continue
    else:doc.add_paragraph(line[2:] if line.startswith('- ') else line,style='List Bullet' if line.startswith('- ') else None)
    i+=1
doc.save(out/'Connectome_Research_Record_2026-09-22_Design_Audit.docx')
paths=[old,source,Path(__file__),root/'docs/pcdr/COMPOSITION_AUDIT_20260922.md',root/'docs/pcdr/CCR_READINESS.md',root/'docs/pcdr/FOUR_WEEK_WORKPLAN.md',root/'docs/pcdr/LAB_NOTEBOOK.md',root/'docs/pcdr/PROCESS_DETAILS.md',root/'results/pcdr/composition_audit_20260922/verification.json',root/'results/pcdr/composition_audit_20260922/motor_witness/summary.json',root/'results/pcdr/composition_audit_20260922/summary.json',root/'results/pcdr/seed_replication_20260921/amendment.json',root/'results/pcdr/seed_replication_20260921/analysis/results.json',root/'results/pcdr/seed_replication_20260921/completion_audit.json']
(out/'source_manifest.json').write_text(json.dumps({'created_utc':datetime.now(timezone.utc).isoformat(),
    'sha256':{str(p.relative_to(root)):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},
    'note':'Cover status updated; original chapters preserved; dated Chapter 13 appended.'},indent=2),encoding='utf-8')
print(out)
