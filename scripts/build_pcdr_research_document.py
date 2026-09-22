"""Consolidate the research record into an editable, navigable Word document."""
from pathlib import Path
import hashlib
import json
import re
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'exports/pcdr_research_record_20260921_completed'
OUT.mkdir(parents=True,exist_ok=True)
DOCS=ROOT/'docs/pcdr'
PILOT=ROOT/'results/pcdr/corrected_20260919'
def read(p): return json.loads(Path(p).read_text(encoding='utf-8'))
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
capture=OUT/'source_snapshot.json'
sources=[DOCS/n for n in ['PLAN.md','PROCESS_DETAILS.md','CODE_GUIDE.md','LAB_NOTEBOOK.md','NEXT_STEPS_20260921.md']]
sources += [PILOT/'report/summary.json',PILOT/'analysis/selection.json',PILOT/'input_audit.json',
            ROOT/'results/pcdr/ccr_singles_20260919/local_status.json',ROOT/'results/pcdr/local_followthrough_20260920/status.json']
sources += [ROOT/'results/pcdr/ccr_singles_20260919/FINDINGS.md', ROOT/'results/pcdr/local_followthrough_20260920/completion_audit.json', ROOT/'results/pcdr/local_followthrough_20260920/FINAL_HANDOFF.md']
sources += [ROOT/'results/pcdr/exploratory80_20260921/FINAL_HANDOFF.md', ROOT/'results/pcdr/exploratory80_20260921/support_audit/FINDINGS.md', ROOT/'results/pcdr/exploratory80_20260921/completion_record.json']
if not capture.exists():
    snapshot={'captured_utc':datetime.now(timezone.utc).isoformat(),
              'sources':{p.relative_to(ROOT).as_posix():{'sha256':sha(p),'content':p.read_text(encoding='utf-8')} for p in sources}}
    capture.write_text(json.dumps(snapshot,indent=2),encoding='utf-8')
snapshot=read(capture)
def source(p): return snapshot['sources'][Path(p).relative_to(ROOT).as_posix()]['content']
summary=json.loads(source(PILOT/'report/summary.json'))
status=json.loads(source(ROOT/'results/pcdr/ccr_singles_20260919/local_status.json'))
follow=json.loads(source(ROOT/'results/pcdr/local_followthrough_20260920/status.json'))
selection=json.loads(source(PILOT/'analysis/selection.json'))
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.72)
sec.left_margin=sec.right_margin=Inches(.78)
sec.header_distance=sec.footer_distance=Inches(.32)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',25),('Subtitle',12),('Heading 1',19),('Heading 2',14),('Heading 3',11.5)]:
    s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(12 if name.startswith('Heading') else 0)
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.keep_with_next=True
for name in ['List Bullet','List Bullet 2']:
    doc.styles[name].paragraph_format.space_after=Pt(4)
    doc.styles[name].paragraph_format.line_spacing=1.06
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Connectome research record  |  ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='Connectome perturbation research record'
doc.core_properties.subject='Research protocol, eigencircuit guide, verified pilot results and dated implementation record'
doc.core_properties.author='Research support prepared with OpenAI Codex'
doc.core_properties.keywords='FlyWire v630; eigencircuits; Shiu LIF; research notebook'
# The bundled Word template carries a title rule; remove paragraph borders.
for container in [doc.element,doc.styles.element]:
    for border in list(container.xpath('.//w:pBdr')):
        border.getparent().remove(border)
bookmarks=0

def link(p,label,url):
    element=OxmlElement('w:hyperlink')
    element.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'24577A');props.append(color)
    run.append(props);text=OxmlElement('w:t');text.text=label;run.append(text);element.append(run);p._p.append(element)

def inline(p,text,origin=None):
    text=text.replace('**','')
    pattern=r'(\[[^\]]+\]\([^)]+\)|`[^`]+`)'
    for piece in re.split(pattern,text):
        match=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',piece)
        if match:
            label,url=match.groups()
            if not url.startswith(('https://','http://')):
                target=((origin.parent if origin else ROOT)/url).resolve()
                # Links point to the local evidence; labels stay readable in a shared copy.
                url=target.as_uri()
            link(p,label,url)
        elif piece.startswith('`') and piece.endswith('`'):
            value=piece[1:-1].replace('\\','/')
            if len(value)>45: value=value.replace('/','/\u200b').replace('_','_\u200b')
            r=p.add_run(value);r.font.name='Consolas';r.font.size=Pt(9.5)
        else:p.add_run(piece)

def para(text,style=None,origin=None):
    p=doc.add_paragraph(style=style);inline(p,text,origin);return p

def heading(text,level=1):
    clean=re.sub(r'[^\w\s]',' ',text).replace('_',' ')
    clean=re.sub(r'\s+',' ',clean).strip()
    return doc.add_heading(clean,level)

def chapter(text):
    heading(text,1).paragraph_format.page_break_before=True

def table(rows,widths=None):
    if not rows:return
    count=len(rows[0]);t=doc.add_table(rows=1,cols=count);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:
        if count==8: widths=[.8,1.65,.55,.75,.85,.9,.65,.79]
        elif count==2: widths=[2.0,4.94]
        elif count==3:widths=[1.6,2.6,2.74]
        else:widths=[6.94/count]*count
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');margins.append(x)
            tcpr.append(margins)
            if i==0 or i%2==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E4EDF2' if i==0 else 'F6F7F8');tcpr.append(fill)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.02
            p.paragraph_format.keep_with_next = i == 0
            if re.fullmatch(r'[-+\d.,%– ]+',str(text)):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,str(text).replace('_',' ') if count==8 and j==0 else str(text),DOCS/'CODE_GUIDE.md')
            for r in p.runs:r.font.size=Pt(9 if count==8 else 10);r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def markdown(content,origin,skip_title=True):
    lines=content.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('```'):
            i+=1;code=[]
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            for value in code:
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2)
                value=value.replace('\\','/').replace('/','/\u200b')
                r=p.add_run(value);r.font.name='Consolas';r.font.size=Pt(9)
            i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                entries=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch('[-: ]+',x) for x in entries):rows.append(entries)
                i+=1
            table(rows);continue
        h=re.match(r'^(#{1,4})\s+(.*)',line)
        if h:
            if not(skip_title and len(h[1])==1 and i==0):heading(h[2],max(2,min(3,len(h[1]))))
        elif line.startswith('- '):para(line[2:],'List Bullet',origin)
        elif line.startswith('> '):para(line[2:],origin=origin)
        else:para(line,origin=origin)
        i+=1

def figure(path,caption,width=6.6):
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(width))
    pic=p._p.xpath('.//wp:docPr')
    if pic:pic[0].set('descr',caption)
    p=para(caption);p.paragraph_format.space_after=Pt(9)
    for r in p.runs:r.font.size=Pt(9.5)

def mr(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
def sub(base,index):
    s=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(base));i=OxmlElement('m:sub');i.append(mr(index));s.extend([e,i]);return s
def fraction(top,bottom):
    f=OxmlElement('m:f')
    for tag,text in [('num',top),('den',bottom)]:
        el=OxmlElement('m:'+tag)
        for node in ([mr(text)] if isinstance(text,str) else text):el.append(node)
        f.append(el)
    return f
def magnitude_sum(lower,upper=None):
    n=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
    for tag,value in [('chr','∑'),('limLoc','subSup'),('supHide','1' if upper is None else '0')]:
        x=OxmlElement('m:'+tag);x.set(qn('m:val'),value);pr.append(x)
    n.append(pr)
    for tag,value in [('sub',lower),('sup',upper or '')]:
        x=OxmlElement('m:'+tag);x.append(mr(value));n.append(x)
    e=OxmlElement('m:e');e.extend([mr('|'),sub('d','i'),mr('|')]);n.append(e)
    return n
def equation(elements):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    math=OxmlElement('m:oMath')
    for e in elements:math.append(mr(e) if isinstance(e,str) else e)
    p._p.append(math)

doc.add_paragraph('Connectome perturbation research record','Title')
doc.add_paragraph('FlyWire v630 eigencircuits and sugar driven LIF simulations','Subtitle')
para('Research support record • 18–21 September 2026')
para('This document brings together the research question, paper annotations, implemented protocol, verified local pilot, eigencircuit mathematics, code guide, process decisions and dated lab notebook. It is intended to make the work understandable and reproducible for the student and research advisers.')
para('The 720-trial individual-cell study is complete. One selected-cell motor effect passed the declared secondary correction. The first 20 stable eigen-mode supports contained no baseline-spiking cells, so none met the recruitment requirement. A later exploratory 80-pair search selected a recruited support, but strict matching accepted no controls in 50,000 proposals. No eigen-set lesion comparison was run; the primary hypothesis remains untested.')
para('Primary endpoints: within-support response A and footprint concentration F jointly. MN9 and total/per-neuron motor ΔHz are secondary simulation readouts. MN9 is not a direct measurement of feeding behavior.')
para('Prepared with OpenAI Codex assistance. The notebook records that support honestly. This is supporting research documentation, not a student-authored STS submission report.')
para('Status snapshot: '+snapshot['captured_utc']+' UTC. Later progress is not silently folded into this edition.')
heading('How to use this document',2)
para('Read Chapters 1–2 for the current position, completed findings and next-step recommendation. Chapters 3–4 cover the initial pilot and eigencircuit concepts. Chapters 5–7 explain the protocol, code and decisions. Chapter 8 preserves the dated notebook; Chapter 9 locates the evidence.')
heading('Contents',2)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u')
doc.add_paragraph()._p.append(field)

chapter('1 Research question and current position')
para('Do Pospisil-style 75%-power eigen-sets of Shiu’s signed connectivity matrix W, computed de novo on FlyWire v630, predict output-lesion footprints in the sugar-evoked LIF model after controlling for degree, strong-synapse mass and recruitment?')
table([['Account','Prediction examined','Comparison needed'],['P  Eigencircuit organization','Eigen-defined supports add predictive localization.','Focal eigen-set against matched lesion sets.'],['D  Degree and strength','Connectivity differences explain the apparent advantage.','Incoming/outgoing degree and absolute strength matching.'],['C  Strong connections','Concentrated strong synaptic mass explains the apparent advantage.','Additional strong-mass matching; separate pruning sensitivity.'],['R  Recruitment','Effects depend on participation under sugar drive.','Baseline recruitment/rate matching and silent controls.']])
para('These accounts can coexist. A nonsignificant result does not establish D, C or R. This study tests a prediction inside a particular nonlinear simulation, not an in-vivo effectome or complete dynamical independence.')
heading('Why localization is primary',2)
para('A downstream motor change can be large even when the response is broad and unrelated to support localization. A and F directly summarize the predicted footprint. F alone can be high for a tiny response; A alone can be high during widespread changes. The primary claim therefore requires both. The user delegated the endpoint choice on 20 September, and the approved localization-primary design was retained before any mode-lesion result.')
table([['Stage','Evidence status at the snapshot'],['Corrected local pilot','Complete: five baselines, one replay, one no-input control, first E and I lesions.'],['Frozen individual-cell study',f"{status['completed']} of {status['total']} trials complete; controller status {status['status']}. Complete thirty-seed study; full findings in Chapter 2."],['Local mode pilot',f"Follow-through status {follow['status']}. No eligible support; controls and mode lesions were not run."],['Exploratory 80-pair extension','Selected a 51-cell support; 0/5 strict controls after 50,000 proposals. Stopped before lesions.'],['Full mode confirmation','Not run. Requires 199 strict controls and 30 seeds; 6,030 trials with baseline.'],['Sensitivities and JO extension','Deferred; see the protocol and reasons in Chapters 5 and 7.']])
para('The original study stopped at eligibility; the later exploratory extension stopped at matching. All controllers exited. Monitoring is paused after handoff. The Monday compute allowance was a maximum budget, not a requirement to run extra stages after a stopping rule.')

chapter('2 Completed study and recommended next steps')
markdown(source(DOCS/'NEXT_STEPS_20260921.md'),DOCS/'NEXT_STEPS_20260921.md')
heading('Completed extension handoff',2)
markdown(source(ROOT/'results/pcdr/exploratory80_20260921/FINAL_HANDOFF.md'),ROOT/'results/pcdr/exploratory80_20260921/FINAL_HANDOFF.md')
heading('Saved support annotation results',2)
markdown(source(ROOT/'results/pcdr/exploratory80_20260921/support_audit/FINDINGS.md'),ROOT/'results/pcdr/exploratory80_20260921/support_audit/FINDINGS.md')
heading('All individual-cell results and secondary tests',2)
markdown(source(ROOT/'results/pcdr/ccr_singles_20260919/FINDINGS.md'),ROOT/'results/pcdr/ccr_singles_20260919/FINDINGS.md')

chapter('3 Verified local pilot results')
para('The results in this chapter come from corrected_20260919 and seeds 630101–630105. The same seed was repeated only as a reproducibility check. Historical runs with missing seed provenance, and the failed original input-pairing comparison, are excluded from these paired conclusions.')
table([['Seed','Recruited non-input cells','MN9 Hz','Motor total Hz','Motor mean Hz']]+[[str(r['seed']),str(r['recruited_noninput']),f"{r['mn9_hz']:.0f}",f"{r['motor_total_hz']:.0f}",f"{r['motor_mean_hz']:.2f}"] for r in summary['baseline']], [1.08,1.55,1.08,1.32,1.32])
para('Same-seed baseline replay produced identical spikes and scheduled inputs. The undriven control produced zero spikes. Baseline simulation/output time was 14.2–16.7 seconds, with peak process memory about 2.84 GB. These are local measurements, not guaranteed cluster performance.')
figure(PILOT/'report/baseline_rates.png','Figure 1. Five independent sugar selection trials. These replicates are simulation seeds, not animals.')
heading('Individual output lesions',2)
table([['Target sign and root ID','MN9 ΔHz','Motor total ΔHz','Motor mean ΔHz']]+[[r['role']+'\n'+r['support_ids'][0],f"{r['mn9_delta_hz']:.0f}",f"{r['motor_total_delta_hz']:.0f}",f"{r['motor_mean_delta_hz']:.3f}"] for r in summary['lesions']],[2.65,1.1,1.45,1.4])
para('Each lesion used one pair with seed 630101. The E target had A = 2 Hz and F = 0.003565; the I target had A = 0 Hz and F = 0 despite nonzero off-target effects. Output silencing does not force a neuron’s own rate to zero. These individual-cell values do not test the eigen-set hypothesis or establish an average E/I effect.')
figure(PILOT/'report/excitatory/footprint.png','Figure 2. Largest rate changes after the first E-cell output lesion, one paired seed. Full all-neuron footprints are saved; the chart shows only fifteen cells.',6.0)
figure(PILOT/'report/inhibitory/footprint.png','Figure 3. Largest rate changes after the first I-cell output lesion, one paired seed. A negative downstream response is possible in a recurrent network.',6.0)
heading('Correlation groups and uncertainty',2)
table([['Bins and pooled count threshold','Cells','Groups','Largest group']]+[[r['analysis'].replace('bins_','').replace('_min_',' / ≥').replace('_',' '),str(r['n_neurons']),str(r['n_clusters']),str(r['largest_cluster'])] for r in summary['correlations']], [3.0,1.1,1.1,1.5])
figure(PILOT/'analysis/bins_10ms_min_1.png','Figure 4. Original and average-linkage order for the 387-cell matrix. Reordering exposes patterns; it does not create evidence of a causal circuit.')
para('At 10 ms, split-trial adjusted Rand index was 0.528 across 366 common cells. Observed within-group correlation was 0.419. Reclustering 99 circular-shift surrogates gave median 0.407; 99 trial-shuffle surrogates gave 0.399. Because each surrogate has its own groups, these summaries do not test preservation of the original memberships. The groups are descriptive sampling strata, not established independent circuits.')
para('Evidence: report/summary.json, per-lesion summary.json and footprint.parquet, analysis/correlation_summary.json, analysis/surrogates.parquet, replay_check.json and the no-input manifest, all under results/pcdr/corrected_20260919. The source snapshot and Chapter 9 identify these records.')

chapter('4 A specific guide to eigencircuits')
heading('From directed edges to a matrix',2)
para('W stores signed synapse counts with postsynaptic cells in rows and presynaptic cells in columns. An entry in row B, column A describes A’s outgoing connection to B. For A → B with +2 and B → C with −3, a vector active only at A gives a positive matrix contribution at B. A vector active only at B gives a negative contribution at C. Integer matrix indices are coordinates in this exact neuron order; they are not root IDs.')
heading('What an eigenvector means here',2)
equation(['Wv = λv'])
para('Multiplication by W changes an eigenvector’s scale by λ while preserving its pattern in the algebraic sense. Each entry of v is a loading on one modeled cell. Large loading magnitude means a cell contributes strongly to that structural pattern; it is not a measured firing rate, causal effect, or probability that the cell controls behavior.')
para('The LIF model adds membrane and synaptic state, thresholds, resets, delays, refractory rules and external drive. Its response need not follow W’s eigenvectors. That disagreement is exactly why structural supports must be perturbed in the simulation.')
heading('How the 75 percent support is selected',2)
para('Take the squared magnitude of every complex loading, normalize by their sum, sort from largest to smallest, and stop at the first cumulative sum reaching 0.75. The selected cells form the support. The threshold concerns loading power, not 75 percent of cells, synapses or spikes.')
table([['Illustrative cell','Normalized loading power','Cumulative power','Selected'],['A','0.60','0.60','Yes'],['B','0.20','0.80','Yes'],['C','0.10','0.90','No'],['D','0.10','1.00','No']],[1.45,1.9,1.8,1.55])
para('This is an illustrative loading-power example, not a computed fly mode. The support contains A and B because A alone is insufficient and A plus B reaches 80 percent. Exact ties use stable neuron-index order so the numerical convention is reproducible.')
heading('Complex modes and repeated eigenvalues',2)
para('A directed signed real matrix can have complex eigenvalues. For loadings [1, i, 1, i], all four squared magnitudes equal one; three cells carry 75 percent. Taking only the real part discards two loadings and changes the support. Complex-conjugate eigenvectors have the same loading magnitudes and count as one paired mode here. Equal absolute eigenvalues alone are not a conjugate pair: +2 and −2 are different real modes.')
para('Repeated or nearly repeated eigenvalues create another issue: more than one basis can describe the same subspace. A solver may return different valid vectors. This protocol does not call an unstable arbitrary basis a unique circuit. It uses two seeded solver starts, residual checks, complete conjugate pairing, exact support agreement and phase-invariant vector overlap.')
heading('How the fly support is chosen',2)
para('Compute the leading 40 eigenpairs of the exact signed v630 matrix; increase to 80 only if needed for twenty complete distinct real/conjugate modes. Among the first twenty stable complete modes, require at least ten support cells with five or more pooled selection spikes. Exclude supports containing directly stimulated sugar inputs. Select greatest loading power on baseline-recruited cells, breaking ties by eigenvalue magnitude and then recorded rank. An unsuccessful selection is reported without replacing sugar or changing thresholds.')
para('This produces one baseline-selected sugar-context mode, not a random sample of all eigencircuits. The v783 paper’s eigenvectors are never imported into the v630 simulation. Multiplying all signed counts by the same global voltage scale changes eigenvalues, not eigenvectors.')
heading('From a support to a lesion footprint',2)
para('For each seed, simulate the baseline and output lesion using the identical scheduled sensory events. For each neuron, subtract baseline Hz from lesion Hz. Average signed differences across seeds before taking absolute values. Let d denote that mean signed ΔHz vector and S the support.')
equation([sub('A','S'),' = ',fraction([magnitude_sum('i∈S')],'|S|')])
equation([sub('F','S'),' = ',fraction([magnitude_sum('i∈S')],[magnitude_sum('i=1','N')])])
para('N is the number of modeled neurons; the denominator of F includes all of them.')
para('Example: mean changes [−2, +1, +1, 0] Hz and support {A, B} give inside absolute change 3 Hz, total absolute change 4 Hz, A = 1.5 Hz and F = 0.75. The signed mean inside is −0.5 Hz. These answer different questions and are all retained. If every change is zero, F is undefined.')
para('For one cell with seed changes [−2, +2], its mean signed change is zero; the primary absolute mean is therefore zero. Averaging the absolute changes instead would give 2 Hz, a different estimand. The code deliberately implements the former.')
heading('Why matching is necessary',2)
para('A support might have unusual degree, total strength, strong outgoing mass or recruitment even without special eigen-organization. Compare it with equally sized, distinct sets matching those measured properties. Match log-transformed continuous features and exact sign/recruitment counts. Each control is scored on its own support. Set overlaps are recorded; matching is conditional on this sampling model, not randomized biological assignment.')
heading('What the result could mean',2)
para('Superiority on both A and F would support this localization prediction in the tested model and sugar context, conditional on the controls. An MN9 change alone supports neither localization nor feeding behavior. Failure to find superiority may reflect a weak prediction, limited power, insufficient recruitment, difficult matching or model assumptions. It does not automatically prove one alternative account.')

chapter('5 Research plan and paper annotations')
para('The working protocol below retains its original dates and explicit amendments. Later local-budget amendments supersede earlier compute windows; they do not erase those earlier records.')
markdown(source(DOCS/'PLAN.md'),DOCS/'PLAN.md')

chapter('6 Function guide and reproduction commands')
markdown(source(DOCS/'CODE_GUIDE.md'),DOCS/'CODE_GUIDE.md')

chapter('7 Process explanations and deferred work')
markdown(source(DOCS/'PROCESS_DETAILS.md'),DOCS/'PROCESS_DETAILS.md')

chapter('8 Dated lab notebook')
para('Entries retain their dates and sequence. A statement that a stage had not yet run describes its status at that entry, not an automatic claim about the export date. Later amendments and the timestamped status in Chapter 1 resolve changes.')
markdown(source(DOCS/'LAB_NOTEBOOK.md'),DOCS/'LAB_NOTEBOOK.md')

chapter('9 Evidence index and reference tools')
heading('Frozen individual targets',2)
para('Selection seed 630500. These cells were chosen from baseline information before their lesion effects were viewed. This table records every selected target, including the three auxiliary controls.')
table([['Role','Root ID','Baseline Hz','Cluster']]+[[c['role'].replace('_',' '),c['root_id'],f"{c['baseline_hz']:.1f}",str(c.get('cluster','—'))] for c in selection['cells']], [1.65,2.55,1.25,1.25])
heading('Evidence locations',2)
table([['Record','Repository relative location'],['Completed pilot and figures','results/pcdr/corrected_20260919/report/'],['All pilot artifacts and hashes','results/pcdr/corrected_20260919/research_record.json'],['Pilot commands and test evidence','results/pcdr/corrected_20260919/commands.json and test_record.json'],['Input audit','results/pcdr/corrected_20260919/input_audit.json'],['Frozen baselines and target selection','results/pcdr/corrected_20260919/analysis/'],['Completed single-cell study','results/pcdr/ccr_singles_20260919/'],['Mode follow-through and execution record','results/pcdr/local_followthrough_20260920/'],['Endpoint and memory decisions','results/pcdr/local_followthrough_20260920/decision_20260920.json'],['Prior excluded input pairing','results/pcdr/corrected_20260919/prior_pair_failure.json'],['This edition source and export hashes','exports/pcdr_research_record_20260921_completed/source_snapshot.json and export_manifest.json']])
para('The raw spike/Parquet datasets and code archives remain separate research evidence. This document embeds the key results and explanations, not every raw data row. Local evidence hyperlinks require this repository on disk; paper hyperlinks are public. Repository-relative paths remain readable when sharing the document.')
heading('Environment used for simulation',2)
markdown('```text\n'+(DOCS/'environment-local.txt').read_text()+'\n```',DOCS/'environment-local.txt',False)
heading('Short glossary',2)
table([['Term','Meaning in this study'],['Connectome','The reconstructed neurons and their chemical synaptic connections.'],['Effectome','Causal influences among neurons; not directly measured by this simulation.'],['LIF','Leaky integrate-and-fire model with voltage dynamics and a spike threshold/reset.'],['Eigen-set or support','Cells carrying the chosen 75 percent of squared complex eigenvector loading.'],['Recruitment','A cell fired under the specified baseline drive; criteria depend on the declared analysis.'],['Output lesion','Set outgoing synaptic weights to zero while preserving incoming connections and the cell.'],['Paired seed','Baseline and lesion share the same scheduled external input realization.'],['A','Mean absolute mean paired rate change within the support, measured in Hz.'],['F','Fraction of whole-network absolute mean change lying inside the support.'],['SMD','Standardized mean difference used to assess feature balance.'],['BH FDR','Benjamini–Hochberg adjustment across the declared family of secondary tests.'],['CCR','Planned cluster execution venue; presently unavailable to the user.'],['Manifest','Machine-readable record of settings, completion, versions, seeds and hashes.']])
heading('Scope of the supplied background',2)
para('The initial pasted proposal and attached design files were treated as proposals and reading leads. Their instructions did not supersede the user’s approved plan. Unsupported statements about complete circuit independence, global metrics being settled, universal strong-edge thresholds or simulations never having been tested were not adopted. Paper annotations in Chapter 5 distinguish verified support from unresolved correspondence. James and Dr Muldoon’s sequence is recorded as user-reported meeting context, not a meeting observed by the assistant.')
heading('Export record',2)
para('Earlier editions remain preserved. source_snapshot.json records this edition’s source text; export_manifest.json records hashes and render checks. The builder is scripts/build_pcdr_research_document.py.')
path=OUT/'Connectome_Research_Record_2026-09-21_Completed.docx'
doc.save(path)
print(path)
print('Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables))
